# service/exporter.py
import os
import time
from pathlib import Path
from datetime import datetime
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import xmind
from xmind.core.markerref import MarkerId
from config import EXPORT_ROOT
from utils.logger import get_logger

logger = get_logger("exporter")

# ---------- 中文自动字体检测 ----------
def get_chinese_font():
    possible = [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simsun.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/System/Library/Fonts/PingFang.ttc"
    ]
    for p in possible:
        if Path(p).exists():
            return p
    raise RuntimeError("未找到中文字体，请安装或配置")

FONT_PATH = get_chinese_font()
pdfmetrics.registerFont(TTFont("ChineseFont", FONT_PATH))

# ---------- DOCX 导出 ----------
def export_docx(title, transcript, summary, todos, dialogue_text="", en_dialogue="", filename=None):
    os.makedirs(EXPORT_ROOT, exist_ok=True)
    if not filename:
        filename = f"{title}_{int(time.time())}.docx"
    path = Path(EXPORT_ROOT) / filename

    doc = Document()
    doc.add_heading(title, 0)
    if dialogue_text:
        doc.add_heading("完整对话（中文）", level=1)
        doc.add_paragraph(dialogue_text)
    if en_dialogue:
        doc.add_heading("英文翻译", level=1)
        doc.add_paragraph(en_dialogue)
    doc.add_heading("会议纪要", level=1)
    doc.add_paragraph(summary)
    doc.add_heading("待办事项", level=1)
    doc.add_paragraph(todos)
    doc.save(str(path))
    logger.info(f"DOCX导出: {path}")
    return str(path)

# ---------- PDF 导出 ----------
def export_pdf(title, summary, todos, dialogue_text="", en_dialogue="", filename=None):
    os.makedirs(EXPORT_ROOT, exist_ok=True)
    if not filename:
        filename = f"{title}_{int(time.time())}.pdf"
    path = Path(EXPORT_ROOT) / filename

    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    margin = 50
    y = height - margin
    c.setFont("ChineseFont", 16)
    c.drawString(margin, y, title)
    y -= 30

    def draw_text(text, font_size=11):
        nonlocal y
        c.setFont("ChineseFont", font_size)
        for line in text.split('\n'):
            while len(line) > 90:
                c.drawString(margin, y, line[:90])
                y -= 18
                line = line[90:]
            c.drawString(margin, y, line)
            y -= 18
            if y < margin:
                c.showPage()
                y = height - margin
                c.setFont("ChineseFont", font_size)

    if dialogue_text:
        draw_text("【中文对话】")
        draw_text(dialogue_text)
    if en_dialogue:
        draw_text("【英文翻译】")
        draw_text(en_dialogue)
    draw_text("【会议纪要】")
    draw_text(summary)
    draw_text("【待办事项】")
    draw_text(todos)
    c.save()
    logger.info(f"PDF导出: {path}")
    return str(path)

# ---------- PPT 导出（全新美化版） ----------
def export_pptx(title, summary, todos, dialogue_text="", filename=None):
    os.makedirs(EXPORT_ROOT, exist_ok=True)
    if not filename:
        filename = f"{title}_{int(time.time())}.pptx"
    path = Path(EXPORT_ROOT) / filename

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. 标题页
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"生成日期：{datetime.now().strftime('%Y-%m-%d %H:%M')}"

    # 2. 会议摘要
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "会议摘要"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = summary
    tf.word_wrap = True
    font_size = 14
    while len(summary) > 200 and font_size > 10:
        font_size -= 1
    for paragraph in tf.paragraphs:
        paragraph.font.size = Pt(font_size)
        paragraph.font.name = "微软雅黑"
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # 3. 待办事项（项目符号列表）
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "待办任务清单"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.word_wrap = True

    todo_lines = [line.strip() for line in todos.split('\n') if line.strip()]
    if not todo_lines:
        todo_lines = ["（无待办事项）"]

    for i, line in enumerate(todo_lines):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        if "|" in line:
            parts = line.split("|")
            if len(parts) >= 2:
                p.text = f"• {parts[0].strip()} —— 责任人：{parts[1].strip()}"
            else:
                p.text = f"• {line}"
        else:
            p.text = f"• {line}"
        p.level = 0
        p.font.size = Pt(12)
        p.font.name = "微软雅黑"

    # 4. 对话节选（如果有）
    if dialogue_text:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "对话节选"
        content = slide.placeholders[1]
        tf = content.text_frame
        tf.text = dialogue_text[:600] + ("……" if len(dialogue_text) > 600 else "")
        tf.word_wrap = True
        for paragraph in tf.paragraphs:
            paragraph.font.size = Pt(11)
            paragraph.font.name = "微软雅黑"

    # 5. 页码
    for i, slide in enumerate(prs.slides, 1):
        left = Inches(0.5)
        top = prs.slide_height - Inches(0.5)
        width = Inches(1)
        height = Inches(0.3)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.text = f"第 {i} 页"
        tf.paragraphs[0].font.size = Pt(10)
        tf.paragraphs[0].font.color.rgb = RGBColor(150, 150, 150)
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT

    prs.save(str(path))
    logger.info(f"PPTX导出: {path}")
    return str(path)

# ---------- XMind 导出 ----------
def export_xmind(title, summary, todos, filename=None):
    os.makedirs(EXPORT_ROOT, exist_ok=True)
    if not filename:
        filename = f"{title}_{int(time.time())}.xmind"
    path = Path(EXPORT_ROOT) / filename

    workbook = xmind.load(str(path))
    sheet = workbook.getPrimarySheet()
    root = sheet.getRootTopic()
    root.setTitle(title)

    sum_topic = root.addSubTopic()
    sum_topic.setTitle("会议摘要")
    for sent in summary.split('。'):
        if sent.strip():
            st = sum_topic.addSubTopic()
            st.setTitle(sent.strip() + '。')

    todo_topic = root.addSubTopic()
    todo_topic.setTitle("待办事项")
    for line in todos.split('\n'):
        if line.strip():
            st = todo_topic.addSubTopic()
            st.setTitle(line.strip())
            if '高' in line:
                st.addMarker(MarkerId.priority1)
            elif '中' in line:
                st.addMarker(MarkerId.priority2)

    xmind.save(workbook, str(path))
    logger.info(f"XMind导出: {path}")
    return str(path)